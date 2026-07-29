from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path("matbench_gnn_final_report.pptx")

BLUE = RGBColor(49, 60, 177)
BLUE_DARK = RGBColor(29, 39, 138)
TEXT = RGBColor(25, 25, 25)
TEXT_LIGHT = RGBColor(255, 255, 255)
ACCENT = RGBColor(229, 240, 255)
GRAY = RGBColor(240, 242, 247)
BORDER = RGBColor(190, 198, 219)


RESULTS = [
    ("E1", "CGCNN", 0.302169, 0.530287, 0.884731, 0.806929, 48, 50, 2.08, 0.792),
    ("E2", "ALIGNN", 0.231432, 0.471325, 0.908939, 0.862117, 50, 50, 5.74, 12.329),
    ("E3", "M3GNet", 0.268041, 0.553444, 0.874444, 0.871087, 42, 50, 3.16, 3.147),
    ("E4", "ALIGNN + Virtual Node", 0.220077, 0.474616, 0.907663, 0.883811, 48, 50, 6.18, 12.210),
    ("E5", "ALIGNN + Transformer", 0.237450, 0.491768, 0.900869, 0.876873, 46, 50, 6.83, 9.691),
    ("E6", "ALIGNN + Physics Constraint", 0.221208, 0.453924, 0.915539, 0.864048, 50, 50, 5.81, 12.548),
]


def set_run_font(run, size: int, bold: bool = False, color: RGBColor = TEXT, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bg(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(250, 251, 255)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.75))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.color.rgb = BLUE


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_bg(slide)
    box = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(9.5), Inches(0.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run_font(run, 24, bold=True, color=TEXT_LIGHT)
    p.alignment = PP_ALIGN.LEFT
    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(11.8), Inches(0.45))
        stf = sbox.text_frame
        sp = stf.paragraphs[0]
        srun = sp.add_run()
        srun.text = subtitle
        set_run_font(srun, 14, color=BLUE_DARK)


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str], subtitle: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.35), Inches(11.7), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(10)
        for run in p.runs:
            set_run_font(run, 20)
    return slide


def add_table_slide(prs: Presentation, title: str, columns: list[str], rows: list[list[str]], col_widths: list[float], subtitle: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle)
    left, top, width, height = Inches(0.45), Inches(1.35), Inches(12.3), Inches(5.6)
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), left, top, width, height)
    table = shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for i, heading in enumerate(columns):
        cell = table.cell(0, i)
        cell.text = heading
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                set_run_font(run, 12, bold=True, color=TEXT_LIGHT)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if r % 2 else GRAY
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    set_run_font(run, 11)
    return slide


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(247, 249, 255)

    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.0), Inches(12.2), Inches(1.25))
    banner.fill.solid()
    banner.fill.fore_color.rgb = BLUE
    banner.line.color.rgb = BLUE

    title_box = slide.shapes.add_textbox(Inches(0.85), Inches(1.2), Inches(10.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "晶体带隙预测实验最终汇报"
    set_run_font(run, 28, bold=True, color=TEXT_LIGHT)

    sub_box = slide.shapes.add_textbox(Inches(0.95), Inches(2.7), Inches(8.5), Inches(1.2))
    stf = sub_box.text_frame
    stf.word_wrap = True
    lines = [
        "项目：基于晶体图神经网络的 band gap 预测",
        "数据集：鸿之微 crystal_band_100000",
        "日期：2026年4月17日",
    ]
    for idx, text in enumerate(lines):
        p = stf.paragraphs[0] if idx == 0 else stf.add_paragraph()
        r = p.add_run()
        r.text = text
        set_run_font(r, 20, color=BLUE_DARK)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    add_bullets_slide(
        prs,
        "项目目标",
        [
            "在鸿之微 crystal_band_100000 数据集上完成 6 组晶体 GNN 实验。",
            "统一设定：单卡 RTX 3090、随机种子 42、所有实验强制训练 50 个 epoch。",
            "任务是 band gap 回归预测，指标包括 MAE、RMSE、R²、Metal F1、训练时长和显存。",
        ],
    )

    add_bullets_slide(
        prs,
        "数据集与任务",
        [
            "清洗后有效样本数为 95,497，划分为 train 76,397 / val 9,549 / test 9,551。",
            "主目标是 band_gap 回归，辅助字段包括 cbm、vbm、gap_type、formula、density。",
            "E6 中进一步利用 CBM/VBM 构造辅助监督和物理一致性约束。",
        ],
    )

    exp_rows = [
        ["E1", "CGCNN Baseline", "原子图卷积基线", "完成"],
        ["E2", "ALIGNN Baseline", "原子图 + line graph，建模键角信息", "完成"],
        ["E3", "M3GNet Baseline", "三体几何建模", "完成"],
        ["E4", "ALIGNN + Virtual Node", "在 ALIGNN 上加入全局虚拟节点", "完成"],
        ["E5", "ALIGNN + Transformer", "在 ALIGNN 后加入 2 层 4 头 Transformer", "完成"],
        ["E6", "ALIGNN + Physics Constraint", "在 ALIGNN 上加入多头输出与物理约束", "完成"],
    ]
    add_table_slide(
        prs,
        "实验设计",
        ["编号", "模型", "核心改动", "状态"],
        exp_rows,
        [0.8, 3.0, 6.2, 1.2],
        "E4、E5、E6 是分别独立建立在 ALIGNN 基线上的增强实验，不是逐层叠加。",
    )

    result_rows = []
    for exp, model, mae, rmse, r2, f1, best_epoch, _, _, _ in RESULTS:
        result_rows.append([exp, model, f"{mae:.6f}", f"{rmse:.6f}", f"{r2:.6f}", f"{f1:.6f}", str(best_epoch)])
    add_table_slide(
        prs,
        "最终结果总表",
        ["实验", "模型", "MAE", "RMSE", "R²", "Metal F1", "Best Epoch"],
        result_rows,
        [0.8, 3.1, 1.3, 1.3, 1.2, 1.6, 1.2],
    )

    add_bullets_slide(
        prs,
        "结果解读",
        [
            "基线对比上，ALIGNN 明显优于 CGCNN 和 M3GNet，是当前最强基线。",
            "按 MAE 排序：E4 > E6 > E2 > E5 > E3 > E1。",
            "按 RMSE 和 R² 排序：E6 最好，说明物理约束版整体回归更稳。",
            "按 Metal F1 排序：E4 最好，说明 Virtual Node 对金属/非金属判别更有帮助。",
            "E5 没有超过 E4，也没有超过 E2，说明当前 Transformer 设计收益不稳定。",
        ],
    )

    add_bullets_slide(
        prs,
        "增强实验与 ALIGNN 基线对比",
        [
            "E4 vs E2：MAE 从 0.2314 降到 0.2201，Metal F1 从 0.8621 提升到 0.8838，说明 Virtual Node 有效。",
            "E5 vs E2：MAE 从 0.2314 上升到 0.2375，RMSE 也变差，说明当前 Transformer 版暂未带来收益。",
            "E6 vs E2：MAE 从 0.2314 降到 0.2212，RMSE 从 0.4713 降到 0.4539，R² 从 0.9089 提升到 0.9155。",
        ],
    )

    cost_rows = []
    for exp, model, _, _, _, _, _, epochs, hours, gpu in RESULTS:
        note = ""
        if exp == "E1":
            note = "最轻量"
        elif exp == "E2":
            note = "最强基线"
        elif exp == "E4":
            note = "最优 MAE / F1"
        elif exp == "E5":
            note = "收益不明显"
        elif exp == "E6":
            note = "最优 RMSE / R²"
        cost_rows.append([f"{exp} {model}", str(epochs), f"{hours:.2f}", f"{gpu:.3f}", note])
    add_table_slide(
        prs,
        "训练开销对比",
        ["实验", "训练轮数", "训练时长(h)", "峰值显存(GB)", "备注"],
        cost_rows,
        [4.4, 1.4, 1.7, 1.8, 2.0],
    )

    add_bullets_slide(
        prs,
        "E6 物理约束的含义",
        [
            "总损失由四部分组成：主任务带隙误差、CBM/VBM 辅助误差、band gap = CBM - VBM 的一致性约束、以及带隙非负惩罚。",
            "通俗讲，E6 不只是让模型猜准 band gap，还要求它顺便学会 CBM 和 VBM，并且三个量之间要彼此对得上。",
            "只有在 CBM 或 VBM 明确有效的样本上，才计算辅助损失和物理约束，避免零占位样本污染训练。",
        ],
    )

    add_bullets_slide(
        prs,
        "最终结论",
        [
            "最强基线是 ALIGNN。",
            "最优 MAE 和 Metal F1 来自 E4：ALIGNN + Virtual Node。",
            "最优 RMSE 和 R² 来自 E6：ALIGNN + Physics Constraint。",
            "当前 Transformer 增强版 E5 没有优于 ALIGNN 基线，后续仍需继续调整。",
            "本轮实验最关键的结论是：Virtual Node 和 Physics Constraint 都是有效增强方向。",
        ],
    )

    add_bullets_slide(
        prs,
        "下一步工作",
        [
            "继续针对 E5 调整位置编码、全局 token 组织方式和注意力层深度。",
            "进一步分析 E4 和 E6 的误差分布，找出各自提升最明显的材料类型。",
            "如果后续继续发论文或做正式答辩，可将 E4 和 E6 作为两个主要增强方向展开。",
        ],
    )

    prs.save(str(OUT_PATH))
    print(f"saved {OUT_PATH}")


if __name__ == "__main__":
    main()
